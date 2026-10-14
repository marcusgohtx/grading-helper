from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl

directory = Path(__file__).parent / 'independent-output'
for name in ['standard.xlsx', 'api.xlsx', 'literal-labels.xlsx']:
    with ZipFile(directory / name) as archive:
        for path in archive.namelist():
            if path.endswith(('.xml','.rels')):
                ET.fromstring(archive.read(path))
        assert not any('externalLink' in p or 'vbaProject' in p for p in archive.namelist())

wb=openpyxl.load_workbook(directory/'standard.xlsx')
assert wb.sheetnames==['Students','Grading']
students, grading=wb.worksheets
assert students.max_row==107 and grading.max_column==103
assert students.freeze_panes=='A8' and grading.freeze_panes=='D9'
assert all(not students[f'{col}{row}'].protection.locked for row in range(8,108) for col in ['B','C'])
assert all(not grading.cell(row,col).protection.locked and grading.cell(row,col).fill.patternType is None for row in [11,12,15] for col in range(4,104))
assert all(cell.fill.fgColor.rgb=='FFE2F0D9' and cell.protection.locked for sheet in wb for cells in sheet for cell in cells if cell.data_type=='f')
assert grading['A7'].font.bold and students['B7'].font.bold
assert [grading.cell(row,1).value for row in [18,19,20]]==['A','B','Final total']
for column,source_row in [('D',8),('BA',57),('CY',107)]:
    assert grading[f'{column}7'].value==f'=IF(\'Students\'!B{source_row}="","",\'Students\'!B{source_row})'
    assert grading[f'{column}8'].value==f'=IF(\'Students\'!C{source_row}="","",\'Students\'!C{source_row})'
    assert f'TYPE({column}11)<>1' in grading[f'{column}18'].value
    assert f'TYPE({column}12)<>1' in grading[f'{column}18'].value
validation=grading.data_validations.dataValidation[0]
assert validation.type=='decimal' and validation.formula1=='0' and validation.formula2=='100'
assert validation.errorStyle=='stop'
assert str(validation.sqref)=='D11:CY12 D15:CY15'
assert [str(item.sqref) for item in grading.conditional_formatting]==['D11:CY12','D15:CY15']
literal=openpyxl.load_workbook(directory/'literal-labels.xlsx')['Grading']
for address,value in [('A10','=SUM(A1) & <tag>'),('A11','+Literal "quotes"'),('A13','-Label'),('A14','@Literal')]:
    assert literal[address].value==value and literal[address].data_type=='s'
print('Independent XML/openpyxl checks passed: sheets, row/column mapping, green/protected formulas, unfilled editable inputs, decimal validation, conditional formatting, totals placement, no external links/macros and literal escaped labels.')
