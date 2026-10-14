from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import math
import subprocess
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
workbook_path = OUTPUT / "grading.xlsx"
workbook = openpyxl.load_workbook(workbook_path, data_only=False)
assert workbook.sheetnames == ["Students", "Grading"]
students, grading = workbook.worksheets
assert students.max_row == 107
assert grading.max_column == 103
assert students.freeze_panes == "A8"
assert grading.freeze_panes == "D9"
assert students.protection.sheet and grading.protection.sheet
assert students["B8"].protection.locked is False
assert grading["D11"].protection.locked is False
assert grading["D11"].fill.patternType is None
assert grading["D20"].protection.locked
assert grading["D20"].number_format == "0.00"
assert grading["D7"].value == '=IF(\'Students\'!B8="","",\'Students\'!B8)'
assert grading["BA7"].value == '=IF(\'Students\'!B57="","",\'Students\'!B57)'
assert grading["CY7"].value == '=IF(\'Students\'!B107="","",\'Students\'!B107)'
validation = grading.data_validations.dataValidation[0]
assert validation.type == "decimal" and validation.errorStyle == "stop"
assert validation.formula1 == "0" and validation.formula2 == "100"
assert str(validation.sqref) == "D11:CY12 D15:CY15"
for sheet in workbook:
    for cells in sheet:
        for cell in cells:
            if cell.data_type == "f":
                assert cell.fill.fgColor.rgb == "FFE2F0D9", cell.coordinate
                assert cell.protection.locked
with ZipFile(workbook_path) as archive:
    assert not any("externalLink" in name or "vbaProject" in name for name in archive.namelist())
    for name in archive.namelist():
        if name.endswith(".xml") or name.endswith(".rels"):
            ET.fromstring(archive.read(name))
    for name in ["xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"]:
        tree = ET.fromstring(archive.read(name))
        for cell in tree.findall(".//s:c[s:f]", NS):
            assert cell.find("s:v", NS).text is None
literal = openpyxl.load_workbook(OUTPUT / "literal-labels.xlsx")
assert literal["Grading"]["A10"].value == "=SUM(A1) & <test>"
assert literal["Grading"]["A10"].data_type == "s"
assert literal["Grading"]["A11"].value == "+Literal"
capacity = openpyxl.load_workbook(OUTPUT / "capacity.xlsx")
assert capacity["Grading"].max_column == 103
assert len(capacity["Grading"].conditional_formatting) == 30
assert sum(1 for row in capacity["Grading"] for cell in row if cell.protection.locked is False) == 20_000
print("XLSX structure, cell protection, blank caches, green formulas, numeric validation and literal labels passed.")

# Patch only input cells in a disposable ZIP, preserving the generator's native formulas.
def change_cells(xml, entries):
    tree = ET.fromstring(xml)
    for address, (kind, value) in entries.items():
        cell = tree.find(f".//s:c[@r='{address}']", NS)
        assert cell is not None, address
        for child in list(cell):
            cell.remove(child)
        if kind == "s":
            cell.set("t", "inlineStr")
            inline = ET.SubElement(cell, "{" + NS["s"] + "}is")
            ET.SubElement(inline, "{" + NS["s"] + "}t").text = str(value)
        else:
            cell.set("t", kind)
            ET.SubElement(cell, "{" + NS["s"] + "}v").text = str(value)
    return ET.tostring(tree, encoding="utf-8", xml_declaration=True)

student_entries = {}
grading_entries = {}
cases = [
    ("Weighted", [80, 60, 90], 83.33333333333333),
    ("Zero", [0, 0, 0], 0),
    ("Partial", [80, None, 90], None),
    ("Numeric text", ["80", 60, 90], None),
    ("Negative", [-1, 60, 90], None),
    ("Too high", [101, 60, 90], None),
    ("Decimal", [82.5, 60.25, 90.5], ((82.5 * 2 + 60.25) / 3) * .4 + 90.5 * .6),
    ("Text", ["absent", 60, 90], None),
    (None, [100, 100, 100], None),
    ("Boolean", [True, 60, 90], None),
    ("Error", ["#VALUE!", 60, 90], None),
]
expected = {}
for offset, (name, scores, total) in enumerate(cases):
    source_row = offset + 8
    col = openpyxl.utils.get_column_letter(offset + 4)
    if name is not None:
        student_entries[f"B{source_row}"] = ("s", name)
    for score_row, score in zip([11, 12, 15], scores):
        if score is None:
            continue
        if isinstance(score, bool):
            kind, score = "b", int(score)
        elif score == "#VALUE!":
            kind = "e"
        else:
            kind = "s" if isinstance(score, str) else "n"
        grading_entries[f"{col}{score_row}"] = (kind, score)
    expected[f"{col}20"] = total
for index in [49, 99]:
    col = openpyxl.utils.get_column_letter(index + 4)
    student_entries[f"B{index + 8}"] = ("s", f"Student {index + 1}")
    student_entries[f"C{index + 8}"] = ("s", f"00{index + 1}")
    for score_row in [11, 12, 15]:
        grading_entries[f"{col}{score_row}"] = ("n", 100)
    expected[f"{col}20"] = 100
inputs = OUTPUT / "inputs"
recalculated = OUTPUT / "recalculated"
inputs.mkdir(exist_ok=True)
recalculated.mkdir(exist_ok=True)
scenario_path = inputs / "scenarios.xlsx"
with ZipFile(workbook_path) as source, ZipFile(scenario_path, "w", ZIP_DEFLATED) as target:
    for entry in source.infolist():
        contents = source.read(entry.filename)
        if entry.filename == "xl/worksheets/sheet1.xml":
            contents = change_cells(contents, student_entries)
        elif entry.filename == "xl/worksheets/sheet2.xml":
            contents = change_cells(contents, grading_entries)
        target.writestr(entry.filename, contents)
command = [r"C:\Program Files\LibreOffice\program\soffice.exe", "-env:UserInstallation=" + (OUTPUT / "lo-profile").as_uri(), "--headless", "--convert-to", "xlsx", "--outdir", str(recalculated), str(scenario_path)]
result_path = recalculated / "scenarios.xlsx"
result_path.unlink(missing_ok=True)
result = subprocess.run(command, capture_output=True, text=True, timeout=60, creationflags=subprocess.CREATE_NO_WINDOW)
assert result.returncode == 0, result.stderr
assert result_path.exists(), result.stdout + result.stderr
calculated = openpyxl.load_workbook(result_path, data_only=True)
for address, expected_value in expected.items():
    actual = calculated["Grading"][address].value
    if expected_value is None:
        assert actual in [None, ""], (address, actual)
    else:
        assert isinstance(actual, (int, float)) and math.isclose(actual, expected_value, abs_tol=1e-9), (address, actual, expected_value)
assert calculated["Grading"]["BA7"].value == "Student 50"
assert calculated["Grading"]["CY8"].value == "00100"
assert calculated["Grading"]["D18"].value == 29.3333333333333
assert calculated["Grading"]["D19"].value == 54
print(f"LibreOffice recalculation: {len(expected)} scenarios passed, including missing scores, zero, decimals, pasted text, booleans, errors, out-of-range values and student columns 50 and 100.")
